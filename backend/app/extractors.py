import io
import re
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument


class ExtractionError(ValueError):
    pass


def normalize_text(text: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def extract_pdf(content: bytes) -> str:
    try:
        with pymupdf.open(stream=content, filetype="pdf") as document:
            return normalize_text("\n".join(page.get_text("text") for page in document))
    except (pymupdf.FileDataError, RuntimeError) as exc:
        raise ExtractionError("Invalid PDF document.") from exc


def extract_docx(content: bytes) -> str:
    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError("Invalid DOCX document.") from exc

    pieces = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            pieces.append(" | ".join(cell.text for cell in row.cells))
    return normalize_text("\n".join(pieces))


def extract_text(content: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(content)
    if suffix == ".docx":
        return extract_docx(content)
    raise ExtractionError("Unsupported document type.")
