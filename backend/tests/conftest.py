import io
import os
from collections.abc import Generator

import pymupdf
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import delete

os.environ.setdefault(
    "DATABASE_URL", "postgresql+psycopg://documind:documind@localhost:5434/documind_test"
)

from app.db import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Document  # noqa: E402


@pytest.fixture(autouse=True)
def clean_documents() -> Generator[None, None, None]:
    with SessionLocal.begin() as session:
        session.execute(delete(Document))
    yield


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def pdf_bytes(text: str = "") -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    result = document.tobytes()
    document.close()
    return result


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        doc_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                doc_table.cell(row_index, column_index).text = value
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
